#!/usr/bin/env python3
"""把技术报告的 Markdown 源文件（docs/report/report.md）生成 Word（docx），可选再导出 PDF。

用法（仓库根目录）：
    python3 tools/build_report.py            # → deliverables/MiniDiDi_report.docx
    python3 tools/build_report.py --pdf      # 同时用 LibreOffice 更新目录、导出 PDF，并把目录写进 docx

需要：pandoc（2.19 以上）、python-docx（pip install python-docx）；导出 PDF 还需要 LibreOffice 和它的 Python 接口（uno）。
Windows 上系统 Python 一般没有 uno，分两步：
    python tools/build_report.py
    "C:\\Program Files\\LibreOffice\\program\\python.exe" tools/build_report.py --pdf-only   # 换成本机的安装位置

排版（按常见的中文课程报告格式，可在 make_reference() 里改）：
  A4，上下 2.54 cm、左右 3.17 cm；正文宋体小四（12 pt）、1.5 倍行距、首行缩进 2 字；
  章标题黑体三号居中、每章另起一页；节标题黑体四号；图题、表题宋体五号居中（表题在表上方，图题在图下方）；
  表格为三线表；页眉"综合设计与实践 C · MiniDiDi 技术报告"，页脚居中页码。
报告正文改 Markdown 即可，样式不用手工调。导出过 PDF 的 docx 已经带好目录；只生成 docx 时目录是空的，
在 Word 里右键目录 →"更新域"即可。
"""
import argparse
import copy
import os
import pathlib
import shutil
import subprocess
import sys
import re
import tempfile
import threading
import time
import unicodedata
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'docs', 'report', 'report.md')
OUT_DIR = os.path.join(ROOT, 'deliverables')
DOCX = os.path.join(OUT_DIR, 'MiniDiDi_report.docx')
HEADER_TEXT = '综合设计与实践 C · MiniDiDi 技术报告'

try:
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError:
    docx = None              # --pdf-only 不需要 python-docx（可以用 LibreOffice 自带的 Python 运行）

SONG, HEI, LATIN, MONO = '宋体', '黑体', 'Times New Roman', 'Consolas'


def set_font(style, east=SONG, latin=LATIN, size=None, bold=None):
    f = style.font
    f.name = latin
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        fonts.set(qn(k), latin)
    fonts.set(qn('w:eastAsia'), east)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if fonts.get(qn(k)) is not None:
            del fonts.attrib[qn(k)]
    f.color.rgb = None


def para(style, align=None, before=0, after=0, line=None, first_indent_chars=None, keep_next=None, page_break=None):
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if page_break is not None:
        pf.page_break_before = page_break
    ppr = style.element.get_or_add_pPr()
    if first_indent_chars is not None:
        ind = ppr.get_or_add_ind()          # python-docx 按架构规定的顺序插入元素
        for k in list(ind.attrib):
            del ind.attrib[k]
        # 显式写出（包括 0），否则会继承基础样式"正文"的首行缩进
        ind.set(qn('w:firstLineChars'), str(first_indent_chars * 100))
        ind.set(qn('w:firstLine'), str(first_indent_chars * 240))


class Styles:
    """按显示名查样式。python-docx 对内置样式名有大小写转换，pandoc 模板里的 "Heading 1" 直接按下标查会找不到。"""

    def __init__(self, d):
        self.d = d

    def __getitem__(self, name):
        for s in self.d.styles:
            if s.name == name or s.name.lower() == name.lower():
                return s
        raise KeyError(name)

    def __contains__(self, name):
        try:
            self[name]
            return True
        except KeyError:
            return False


def get_style(d, name, base='Normal'):
    st = Styles(d)
    if name in st:
        return st[name]
    s = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = st[base]
    return s


def three_line_table(style):
    """三线表：表格上下各一条粗线（1.5 磅），表头下一条细线（0.75 磅），没有竖线。"""
    el = style.element
    tblpr = el.find(qn('w:tblPr'))
    if tblpr is None:
        tblpr = OxmlElement('w:tblPr')
        el.append(tblpr)
    old = tblpr.find(qn('w:tblBorders'))
    if old is not None:
        tblpr.remove(old)
    borders = OxmlElement('w:tblBorders')
    later = [tblpr.find(qn(t)) for t in ('w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook')]
    for side, sz in (('top', 12), ('bottom', 12)):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), str(sz))
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        borders.append(b)
    later = [e for e in later if e is not None]
    if later:
        later[0].addprevious(borders)       # tblBorders 必须在 shd / tblLayout / tblCellMar / tblLook 之前
    else:
        tblpr.append(borders)
    for old in el.findall(qn('w:tblStylePr')):
        el.remove(old)
    first = OxmlElement('w:tblStylePr')
    first.set(qn('w:type'), 'firstRow')
    rpr = OxmlElement('w:rPr')
    b = OxmlElement('w:b')
    rpr.append(b)
    first.append(rpr)
    tcpr = OxmlElement('w:tcPr')
    tcb = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), '000000')
    tcb.append(bottom)
    tcpr.append(tcb)
    first.append(tcpr)
    el.append(first)


def field_run(paragraph, instr):
    """在段落里插入一个域（如 PAGE 页码）。"""
    for kind, text in (('begin', None), ('instr', instr), ('separate', None), ('text', '1'), ('end', None)):
        r = OxmlElement('w:r')
        if kind in ('begin', 'separate', 'end'):
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), kind)
            r.append(fc)
        elif kind == 'instr':
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = f' {text} '
            r.append(it)
        else:
            t = OxmlElement('w:t')
            t.text = text
            r.append(t)
        paragraph._p.append(r)


def make_reference(path):
    subprocess.run(['pandoc', '-o', path, '--print-default-data-file', 'reference.docx'], check=True)
    d = docx.Document(path)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)
    sec.header_distance = sec.footer_distance = Cm(1.5)
    pgmar = sec._sectPr.find(qn('w:pgMar'))
    if pgmar is not None and pgmar.get(qn('w:gutter')) is None:
        pgmar.set(qn('w:gutter'), '0')       # 架构要求必须有 gutter 属性

    st = Styles(d)
    for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact', 'Block Text', 'Footnote Text',
                 'Definition', 'Definition Term', 'Bibliography', 'Abstract'):
        set_font(st[name], SONG, LATIN, 12)
    for name in ('Body Text', 'First Paragraph'):
        para(st[name], WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 0, 1.5, first_indent_chars=2)
    # Compact 用于表格单元格和紧凑列表：左对齐，窄列里两端对齐会把"HTTP 8080"之类拉出大空隙。
    # 行距用"最小值"：按倍数算时纯英文的单元格行高比中文小，同一行里文字会高低不齐
    para(st['Compact'], WD_ALIGN_PARAGRAPH.LEFT, 0, 0, None, first_indent_chars=0)
    set_font(st['Compact'], SONG, LATIN, 10.5)
    st['Compact'].paragraph_format.line_spacing = Pt(17)
    st['Compact'].paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    para(st['Normal'], None, 0, 0, 1.5)

    for name, size, align, before, after, pb in (('Heading 1', 16, WD_ALIGN_PARAGRAPH.CENTER, 12, 18, True),
                                                 ('Heading 2', 14, WD_ALIGN_PARAGRAPH.LEFT, 12, 6, False),
                                                 ('Heading 3', 12, WD_ALIGN_PARAGRAPH.LEFT, 6, 3, False)):
        s = st[name]
        set_font(s, HEI, LATIN, size, bold=True)
        para(s, align, before, after, 1.5, first_indent_chars=0, keep_next=True, page_break=pb)

    for name in ('Image Caption', 'Table Caption', 'Caption'):
        s = st[name]
        set_font(s, SONG, LATIN, 10.5, bold=False)
        s.font.italic = False
        para(s, WD_ALIGN_PARAGRAPH.CENTER, 3, 9, 1.25, first_indent_chars=0)
    para(st['Table Caption'], WD_ALIGN_PARAGRAPH.CENTER, 9, 3, 1.25, first_indent_chars=0, keep_next=True)
    for name in ('Figure', 'Captioned Figure'):
        para(st[name], WD_ALIGN_PARAGRAPH.CENTER, 6, 0, 1.0, first_indent_chars=0, keep_next=True)

    s = get_style(d, 'Source Code')
    set_font(s, SONG, MONO, 9)
    para(s, None, 3, 3, 1.0, first_indent_chars=0)

    three_line_table(st['Table'])

    # 封面、摘要标题、参考文献（report.md 里用 ::: {custom-style="…"} 引用）
    for name, east, size, bold, before, after in (('Cover School', HEI, 26, True, 72, 6),
                                                  ('Cover Course', SONG, 18, False, 6, 60),
                                                  ('Cover Title', HEI, 22, True, 24, 12),
                                                  ('Cover Subtitle', SONG, 16, False, 6, 96),
                                                  ('Cover Info', SONG, 14, False, 6, 6)):
        s = get_style(d, name)
        set_font(s, east, LATIN, size, bold)
        para(s, WD_ALIGN_PARAGRAPH.CENTER if name != 'Cover Info' else WD_ALIGN_PARAGRAPH.LEFT, before, after, 1.5,
             first_indent_chars=0)
        if name == 'Cover Info':
            s.paragraph_format.left_indent = Cm(3.5)
    s = st['Abstract Title']
    set_font(s, HEI, LATIN, 16, True)
    para(s, WD_ALIGN_PARAGRAPH.CENTER, 12, 18, 1.5, first_indent_chars=0)
    s = get_style(d, 'Reference')
    set_font(s, SONG, LATIN, 10.5)
    para(s, WD_ALIGN_PARAGRAPH.JUSTIFY, 0, 3, 1.25, first_indent_chars=0)
    s.paragraph_format.left_indent = Cm(0.74)
    s.paragraph_format.first_line_indent = Cm(-0.74)

    # 目录各级的样式
    for lvl, size in ((1, 12), (2, 12)):
        s = get_style(d, f'toc {lvl}')
        set_font(s, HEI if lvl == 1 else SONG, LATIN, size)
        para(s, None, 3 if lvl == 1 else 0, 0, 1.5, first_indent_chars=0)
        s.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.85)

    # 页眉、页脚（pandoc 生成文档时沿用参考文档的页眉页脚）
    hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    hp.text = HEADER_TEXT
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.size = Pt(9)
    pbdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', '000000')):
        bot.set(qn(k), v)
    pbdr.append(bot)
    hppr = hp._p.get_or_add_pPr()
    jc = hppr.find(qn('w:jc'))
    if jc is not None:
        jc.addprevious(pbdr)                 # pBdr 必须在 jc 之前
    else:
        hppr.append(pbdr)
    fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field_run(fp, 'PAGE')

    d.save(path)


def polish_tables(docx_path):
    """表格的两处排版：不超过 12 行的表格不跨页（除最后一行外每行"与下段同页"，这是 Word 的标准做法），
    免得页面底部只剩表头和一行；表格后面紧跟的正文段落加 6 磅段前距，不然文字贴着表格的底线。"""
    d = docx.Document(docx_path)
    for tbl in d.element.body.findall(qn('w:tbl')):
        rows = tbl.findall(qn('w:tr'))
        if len(rows) <= 12:
            for tr in rows[:-1]:
                for p in tr.iter(qn('w:p')):
                    p.get_or_add_pPr().keepNext_val = True
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn('w:p'):
            nxt.get_or_add_pPr().spacing_before = Pt(6)
    d.save(docx_path)


def enable_update_fields(docx_path):
    """让 Word 打开文档时提示"更新域"，目录随之生成。pandoc 会重写 settings.xml，所以在生成之后再加。"""
    tmp = docx_path + '.tmp'
    with zipfile.ZipFile(docx_path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'word/settings.xml' and b'updateFields' not in data:
                tag = b'<w:updateFields w:val="true" />'
                anchor = data.find(b'<w:footnotePr')     # 架构规定 updateFields 在 footnotePr 之前
                if anchor < 0:
                    anchor = data.find(b'</w:settings>')
                data = data[:anchor] + tag + data[anchor:]
            zout.writestr(item, data)
    os.replace(tmp, docx_path)


def display_width(text):
    return sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in text)


TABLE_UNITS = 78        # 正文宽 14.66 cm，表格里 10.5 磅的字：一行约放 78 个半角字符（39 个汉字）
CELL_PAD = 3            # 单元格左右内边距约合 2 个半角字符，再留 1 个余量
ZWSP, NBSP = '\u200b', '\u00a0'      # 零宽空格（允许换行）、不换行空格


def break_long_code(code):
    """很长的驼峰标识符（测试方法名等）在单词边界插入零宽空格，窄列里在那里换行，而不是随便断在某个字母上。"""
    if len(code) <= 20:
        return code
    out, run = [], 0
    for k, c in enumerate(code):
        if run >= 10 and c.isupper() and code[k - 1].islower():
            out.append(ZWSP)
            run = 0
        out.append(c)
        run += 1
    return ''.join(out)


def prepare_table_row(line):
    line = re.sub(r'(?<=\d) (?=\d{3}(?!\d))', NBSP, line)          # "1 978" 这样的分组数字不拆开
    return re.sub(r'`([^`]+)`', lambda m: '`' + break_long_code(m.group(1)) + '`', line)


def widen_tables(md):
    """按内容长度设置表格列宽。

    pandoc 的规则：只要某一行超过 --columns（默认 72）个字符，管道表格就占满页宽，各列宽度按分隔行（|---|）里
    短横线的数量分配。这里按每列内容的显示宽度（中文算 2，过长的按 40 封顶）重写分隔行，源文件里保持 |---| 即可。
    先保证每列放得下表头和这一列里最长的英文单词 / 代码（不然 HeadingEstimator 之类会被拆成两行），
    剩下的页宽先让内容短的列一行放下（如"匈牙利"），再按差多少分给内容长的列。
    """
    def word_width(w):
        return sum(1.4 if c.isupper() else 1.0 for c in w)      # 大写字母比小写宽

    def longest_word(text):
        widths = [0]
        for code in re.findall(r'`([^`]+)`', text):                # 代码用等宽字体，每个字符约 1.1
            widths += [1.1 * len(w) for w in re.split(r'[\s' + ZWSP + ']+', code)]
        text = re.sub(r'`[^`]+`', ' ', text)
        # 英文单词后面紧跟的中文右括号、逗号不能换到下一行开头，算进单词里
        for w in re.findall(r'[A-Za-z0-9_.:/\\' + NBSP + r'-]+[）)，、。]?', text):
            widths.append(word_width(w[:-1]) + 2 if w[-1] in '），、。' else word_width(w))
        return max(widths)

    lines = md.split('\n')
    out = []
    i = 0
    sep = re.compile(r'^\|(\s*:?-+:?\s*\|)+\s*$')
    while i < len(lines):
        if i + 1 < len(lines) and lines[i].startswith('|') and sep.match(lines[i + 1]):
            j = i + 2
            while j < len(lines) and lines[j].startswith('|'):
                j += 1
            for k in [i] + list(range(i + 2, j)):
                lines[k] = prepare_table_row(lines[k])
            rows = [lines[i]] + lines[i + 2:j]
            # 量宽度时去掉加粗、斜体标记（**450 ± 10** 实际只占 450 ± 10 的宽度）
            cells = [[re.sub(r'\*\*|(?<!\\)\*', '', c).strip() for c in r.strip().strip('|').split('|')] for r in rows]
            n = len(cells[0])
            mins, wants = [], []
            for k in range(n):
                col = [r[k] if k < len(r) else '' for r in cells]
                need = max(4, display_width(col[0]), min(30, max(longest_word(c) for c in col))) + CELL_PAD
                mins.append(need)
                wants.append(max(need, min(40, max(display_width(c) for c in col)) + CELL_PAD))
            widths = mins[:]
            left = TABLE_UNITS - sum(mins)
            todo = [k for k in range(n) if wants[k] > mins[k]]
            while left > 0 and todo:
                share = left / len(todo)
                short = [k for k in todo if wants[k] - widths[k] <= share]
                if not short:                                     # 都放不下一行：按各列还差多少分
                    need = sum(wants[k] - widths[k] for k in todo)
                    for k in todo:
                        widths[k] += left * (wants[k] - widths[k]) / need
                    break
                for k in short:
                    left -= wants[k] - widths[k]
                    widths[k] = wants[k]
                    todo.remove(k)
            total = sum(widths)
            dashes = [max(3, round(w * 120 / total)) for w in widths]
            aligns = [c.strip() for c in lines[i + 1].strip().strip('|').split('|')]
            parts = []
            for k in range(n):
                a = aligns[k] if k < len(aligns) else '---'
                left, right = a.startswith(':'), a.endswith(':')
                parts.append((':' if left else '') + '-' * dashes[k] + (':' if right else ''))
            out.append(lines[i])
            out.append('|' + '|'.join(parts) + '|')
            out.extend(lines[i + 2:j])
            i = j
            continue
        out.append(lines[i])
        i += 1
    return '\n'.join(out)


def chinese_quotes(md):
    """把正文里成对的 ASCII 双引号换成中文引号“”。

    pandoc 的智能引号按前一个字符判断开 / 闭，中文字后面的引号会被当成闭引号，结果两边都是”。
    代码块、行内代码、raw openxml 和 ::: {custom-style="…"} 行不处理。
    """
    out, fenced = [], False
    for line in md.split('\n'):
        if line.startswith('```'):
            fenced = not fenced
        if fenced or line.startswith(':::') or line.startswith('<w:'):
            out.append(line)
            continue
        parts = re.split(r'(`[^`]*`)', line)
        for k in range(0, len(parts), 2):
            parts[k] = re.sub(r'"([^"]*)"', '\u201c\\1\u201d', parts[k])
        out.append(''.join(parts))
    return '\n'.join(out)


def to_pdf(docx_path, pdf_path):
    """用 LibreOffice（无界面）打开 docx，更新目录等索引后导出 PDF。

    返回目录的各行 [(标题, 页码)]（供 fill_toc 写回 docx）；失败时返回 None。
    """
    try:
        import uno  # noqa: F401
        from com.sun.star.beans import PropertyValue
    except ImportError:
        print('没有 LibreOffice 的 Python 接口（uno），跳过 PDF；可在 Word 里"另存为 PDF"。'
              'Windows 上可以用 LibreOffice 自带的 Python：<LibreOffice>/program/python.exe tools/build_report.py --pdf-only',
              file=sys.stderr)
        return None
    import uno

    def prop(name, value):
        p = PropertyValue()
        p.Name, p.Value = name, value
        return p

    pipe = f'minididi_{os.getpid()}'
    profile = tempfile.mkdtemp(prefix='lo_profile_')
    # 用 LibreOffice 自带的 Python 运行时 soffice 不一定在 PATH 里，就在同一目录下找
    soffice = shutil.which('soffice') or os.path.join(os.path.dirname(sys.executable), 'soffice')
    # 配置目录必须是合法的 file URL：Windows 上 file://C:\... 不合法，LibreOffice 会误报"bootstrap.ini 已经损坏"
    office = subprocess.Popen([soffice, '--headless', '--invisible', '--nologo', '--norestore',
                               '-env:UserInstallation=' + pathlib.Path(profile).as_uri(), f'--accept=pipe,name={pipe};urp;'],
                              stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
        ctx = None
        for _ in range(180):                 # 新建配置目录后第一次启动可能要几十秒
            try:
                ctx = resolver.resolve(f'uno:pipe,name={pipe};urp;StarOffice.ComponentContext')
                break
            except Exception:
                time.sleep(0.5)
        if ctx is None:
            print('LibreOffice 没有启动', file=sys.stderr)
            return None
        desktop = ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', ctx)
        doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(os.path.abspath(docx_path)), '_blank', 0,
                                           (prop('Hidden', True),))
        indexes = doc.getDocumentIndexes()
        toc = []
        for i in range(indexes.getCount()):
            index = indexes.getByIndex(i)
            index.update()
            if not toc:
                for line in index.getAnchor().getString().splitlines():
                    if '\t' in line:
                        title, page = line.rsplit('\t', 1)
                        if page.strip().isdigit():
                            toc.append((title.replace('\t', ' ').strip(), page.strip()))
        doc.refresh()
        doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(pdf_path)), (prop('FilterName', 'writer_pdf_Export'),))
        # PDF 已经导出。关闭文档偶尔会卡住（LibreOffice 在后台一直排版），最多等 20 秒，然后直接结束进程
        closer = threading.Thread(target=close_quietly, args=(doc,), daemon=True)
        closer.start()
        closer.join(20)
        return toc
    finally:
        stop_office(office)
        shutil.rmtree(profile, ignore_errors=True)
        lock = os.path.join(os.path.dirname(os.path.abspath(docx_path)), '.~lock.' + os.path.basename(docx_path) + '#')
        if os.path.exists(lock):             # LibreOffice 被结束得太快时会留下锁文件
            os.remove(lock)


def close_quietly(doc):
    try:
        doc.close(True)
    except Exception:
        pass                                 # 进程被结束时这里会收到 DisposedException


def stop_office(office):
    if os.name == 'nt':
        # Windows 上启动的 soffice 只是启动器，干活的 soffice.bin 是它的子进程，要连同子进程一起结束
        subprocess.run(['taskkill', '/F', '/T', '/PID', str(office.pid)],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        office.terminate()
        office.wait(timeout=10)
    except Exception:
        office.kill()


def show(path):
    """打印用的路径：在仓库里就显示相对路径（输出到别的盘时 relpath 会报错）。"""
    try:
        return os.path.relpath(path, ROOT)
    except ValueError:
        return os.path.abspath(path)


def heading_levels(md_path):
    """报告源文件里各个一、二级标题的级别：{标题文字: 1 或 2}。"""
    levels = {}
    with open(md_path, encoding='utf-8') as f:
        for line in f:
            m = re.match(r'^(#{1,2}) (.+?)\s*(\{[^}]*\})?\s*$', line)
            if m:
                levels[m.group(2).strip()] = len(m.group(1))
    return levels


def fill_toc(docx_path, entries, levels):
    """把 LibreOffice 算好的目录写进 docx 目录域的"显示结果"。

    报告里的目录是一个 TOC 域，Word 要"更新域"后才显示内容；没点更新、在受保护视图里打开或用 WPS 打开时就是空的。
    写入显示结果后打开即可看到目录；域本身保留，在 Word 里更新域会按 Word 自己的排版重新计算页码。
    只用标准库（zipfile、re），可以在 LibreOffice 自带的 Python 里运行。
    """
    def esc(text):
        return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    with zipfile.ZipFile(docx_path) as z:
        xml = z.read('word/document.xml').decode('utf-8')
        styles = z.read('word/styles.xml').decode('utf-8')
    k = xml.find('TOC \\o')
    if k < 0 or not entries:
        print('没有找到目录域，目录未写入 docx', file=sys.stderr)
        return False
    start = xml.rfind('<w:p>', 0, k)
    # 域可能已经写过一次（跨多个段落）：替换到域结束标记所在段落为止
    end = xml.find('</w:p>', xml.find('w:fldCharType="end"', k)) + len('</w:p>')
    old = xml[start:end]
    sep = '<w:fldChar w:fldCharType="separate"/></w:r>'
    head = old[old.find('<w:r><w:fldChar w:fldCharType="begin"'):old.find(sep) + len(sep)]   # begin、instrText、separate 原样保留
    tail = '<w:r><w:fldChar w:fldCharType="end"/></w:r>'

    def style_id(name, default):
        m = re.search(r'<w:style\b[^>]*w:styleId="([^"]+)"[^>]*>\s*<w:name w:val="' + name + '"', styles)
        return m.group(1) if m else default

    ids = {1: style_id('toc 1', 'TOC1'), 2: style_id('toc 2', 'TOC2')}
    tab_pos = round((21.0 - 2 * 3.17) / 2.54 * 1440)            # 页码右对齐到正文右边界（单位 1/20 磅）
    paras = []
    for i, (title, page) in enumerate(entries):
        level = levels.get(title) or (2 if re.match(r'^[0-9A-Z]+\.\d', title) else 1)
        ppr = (f'<w:pPr><w:pStyle w:val="{ids[level]}"/><w:tabs>'
               f'<w:tab w:val="right" w:leader="dot" w:pos="{tab_pos}"/></w:tabs></w:pPr>')
        runs = (f'<w:r><w:t xml:space="preserve">{esc(title)}</w:t></w:r>'
                f'<w:r><w:tab/></w:r><w:r><w:t>{esc(page)}</w:t></w:r>')
        paras.append('<w:p>' + ppr + (head if i == 0 else '') + runs + (tail if i == len(entries) - 1 else '') + '</w:p>')
    xml = xml[:start] + ''.join(paras) + xml[end:]

    tmp = docx_path + '.tmp'
    with zipfile.ZipFile(docx_path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = xml.encode('utf-8') if item.filename == 'word/document.xml' else zin.read(item.filename)
            zout.writestr(item, data)
    os.replace(tmp, docx_path)
    return True


def main():
    ap = argparse.ArgumentParser(description='生成技术报告 docx（可选 PDF）')
    ap.add_argument('--src', default=SRC)
    ap.add_argument('--out', default=DOCX)
    ap.add_argument('--pdf', action='store_true', help='同时导出 PDF，并把目录写进 docx（需要 LibreOffice）')
    ap.add_argument('--pdf-only', action='store_true',
                    help='不重新生成 docx，只更新目录并导出 PDF（不需要 pandoc / python-docx，可用 LibreOffice 自带的 Python 运行）')
    args = ap.parse_args()
    if args.pdf_only:
        export_pdf(args)
        return
    if shutil.which('pandoc') is None:
        sys.exit('需要 pandoc：https://pandoc.org/installing.html')
    if docx is None:
        sys.exit('需要 python-docx：pip install python-docx')
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        ref = os.path.join(tmp, 'reference.docx')
        make_reference(ref)
        src_dir = os.path.dirname(os.path.abspath(args.src))
        with open(args.src, encoding='utf-8') as f:
            md = widen_tables(chinese_quotes(f.read()))
        prepared = os.path.join(tmp, 'report.md')
        with open(prepared, 'w', encoding='utf-8') as f:
            f.write(md)
        subprocess.run(['pandoc', prepared, '-f', 'markdown-smart', '-t', 'docx',
                        '--reference-doc', ref, '--resource-path', src_dir, '-o', os.path.abspath(args.out)],
                       check=True, cwd=src_dir)
    polish_tables(args.out)
    enable_update_fields(args.out)
    print('已生成', show(args.out))
    if args.pdf:
        export_pdf(args)


def export_pdf(args):
    pdf = os.path.splitext(args.out)[0] + '.pdf'
    toc = to_pdf(args.out, pdf)
    if toc is not None:
        print('已生成', show(pdf))
        if fill_toc(args.out, toc, heading_levels(args.src)):
            print(f'目录（{len(toc)} 条）已写入', show(args.out))


if __name__ == '__main__':
    main()
